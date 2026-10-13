"""One-off check: compare Phase-Summary road start/end dates against each
road sheet's first activity start-date / last activity end-date. Read-only,
reports only -- does not change any file. Run: python3 scripts/diff_dates.py
"""
import os
import re
import datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "Source", "Final Sheet for Claude Project-Phase-1.xlsx")
DMY = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, str):
        m = DMY.match(v.strip())
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            try:
                return datetime.date(y, mo, d)
            except ValueError:
                return None
    return None


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Phase- Summary"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[1]]
    idx = {name: i for i, name in enumerate(hdr)}
    summary = {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        try:
            rn = int(str(r[0]).replace("\xa0", "").strip())
        except (ValueError, TypeError):
            continue
        summary[rn] = (to_date(r[idx["Start date"]]), to_date(r[idx["End Date"]]))

    sheet_by_no = {}
    for name in wb.sheetnames:
        m = re.match(r"Road (\d+)", name)
        if m:
            sheet_by_no[int(m.group(1))] = name

    print(f"{'Road':<5}{'Summary Start':<14}{'Summary End':<14}{'Sheet First Start':<18}{'Sheet Last End':<16}{'Match?'}")
    for rn in sorted(summary):
        s_start, s_end = summary[rn]
        starts, ends = [], []
        if rn in sheet_by_no:
            ws2 = wb[sheet_by_no[rn]]
            started = False
            for row in ws2.iter_rows(values_only=True):
                if not started:
                    if any(str(c).strip() == "Description of Job" for c in row if c):
                        started = True
                    continue
                if len(row) < 4 or not row[1]:
                    continue
                st = to_date(row[4]) if len(row) > 4 else None
                en = to_date(row[5]) if len(row) > 5 else None
                if st:
                    starts.append(st)
                if en:
                    ends.append(en)
        sheet_start = min(starts) if starts else None
        sheet_end = max(ends) if ends else None
        match = "YES" if (sheet_start == s_start and sheet_end == s_end) else "NO"
        print(f"{rn:<5}{str(s_start):<14}{str(s_end):<14}{str(sheet_start):<18}{str(sheet_end):<16}{match}")


if __name__ == "__main__":
    main()
