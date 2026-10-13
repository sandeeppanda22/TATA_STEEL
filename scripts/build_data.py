"""Build the Phase-1 dataset in memory from the source Excel.
Reusable regeneration logic (hard-rebuild workflow). No assumptions:
raw values are carried through; data problems and unmapped activity
names are REPORTED, never silently fixed.

This module writes NO files. scripts/gen_seed.py and scripts/verify_seed.py
import build_dataset() and consume the dict directly, so the pipeline is a
single step: Source/*.xlsx -> db/01_seed.sql (no JSON intermediate).

Run standalone (prints a summary only): python3 scripts/build_data.py
"""
import datetime
import os
import re

import openpyxl

import config

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC_DIR = os.path.join(ROOT, "Source")
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


DMY = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


def iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    return None if v is None else str(v).strip()


def normalize_date(s):
    """Confirmed with user: text dates on Road 3/4 are dd/mm/yyyy (unambiguous,
    day values >12). Convert to ISO. Returns (iso_or_original, was_converted)."""
    if not s or ISO.match(s):
        return s, False
    m = DMY.match(s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime.date(y, mo, d).isoformat(), True
        except ValueError:
            return s, False
    return s, False


def date_issue(s):
    """Return a problem label for a date cell, or None if it looks clean."""
    if not s:
        return "missing"
    if not ISO.match(s):
        return "non-ISO text date"
    return None


def read_summary(wb):
    ws = wb["Phase- Summary"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[1]]
    idx = {name: i for i, name in enumerate(hdr)}
    roads = []
    for r in rows[2:]:
        if r[0] is None:
            continue
        try:
            road_no = int(str(r[0]).replace("\xa0", "").strip())
        except (ValueError, TypeError):
            continue  # 'Total' row or blank
        roads.append({
            "road_no": road_no,
            "scheme_sl_no": iso(r[idx["Scheme SL No"]]),
            "location": iso(r[idx["Location / Road Mark"]]),
            "line_manager": iso(r[idx["LM"]]),
            "vendor": iso(r[idx["Vendor"]]),
            "road_type": iso(r[idx["Road Type"]]),
            "length_m": r[idx["Length"]] or 0,
            "width_m": r[idx["Width"]],
            "drain_m": r[idx["Drain"]] or 0,  # Phase-Summary = absolute truth (confirmed)
            "cost_cr": r[idx["Cost In Cr (Service)"]],
            "start_date": iso(r[idx["Start date"]]),  # Phase-Summary = absolute truth (confirmed)
            "end_date": iso(r[idx["End Date"]]),      # Phase-Summary = absolute truth (confirmed)
        })
    return roads


def read_activities(ws, road_type=None):
    """Pull activity rows from a road sheet (header row has 'Description of Job').
    road_type is threaded through to config.normalize() because bare "RCC"
    resolves differently on a Paver-type road vs an RCC-type road."""
    acts, started = [], False
    for row in ws.iter_rows(values_only=True):
        cells = [iso(c) for c in row]
        if not started:
            if any(c == "Description of Job" for c in cells if c):
                started = True
            continue
        if len(cells) < 4 or not cells[1]:
            continue
        raw = cells[1]
        canon, mapped = config.normalize(raw, road_type=road_type)
        start_raw = cells[4] if len(cells) > 4 else None
        end_raw = cells[5] if len(cells) > 5 else None
        start_norm, start_conv = normalize_date(start_raw)
        end_norm, end_conv = normalize_date(end_raw)
        acts.append({
            "raw_name": raw,
            "activity": canon,
            "mapped": mapped,
            "unit": cells[2],
            "planned_qty": cells[3],
            "start_date": start_norm,
            "end_date": end_norm,
            "start_date_converted": start_conv,
            "end_date_converted": end_conv,
        })
    return acts


def order_key(a):
    seq = config.ACTIVITY_SEQUENCE
    return seq.index(a["activity"]) if a["activity"] in seq else len(seq) + 1


def source_files():
    """Every .xlsx in Source/, ignoring Excel lock files and OS junk."""
    if not os.path.isdir(SRC_DIR):
        return []
    out = []
    for name in sorted(os.listdir(SRC_DIR)):
        if name.startswith("~$") or name.startswith("."):
            continue
        if name.lower().endswith((".xlsx", ".xlsm")):
            out.append(os.path.join(SRC_DIR, name))
    return out


def read_workbook(path, roads_by_no, issues, unmapped):
    """Read one source workbook, appending its roads into roads_by_no."""
    fname = os.path.basename(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Phase- Summary" not in wb.sheetnames:
        issues.append({"road_no": "-", "activity": f"({fname})", "field": "workbook",
                       "value": fname,
                       "issue": "no 'Phase- Summary' sheet - file skipped"})
        return
    sheet_by_no = {}
    for name in wb.sheetnames:
        m = re.match(r"Road (\d+)", name)
        if m:
            sheet_by_no[int(m.group(1))] = name

    for road in read_summary(wb):
        rn = road["road_no"]
        road["source_file"] = fname
        if rn in roads_by_no:
            issues.append({"road_no": rn, "activity": f"({fname})", "field": "road_no",
                           "value": f'also in {roads_by_no[rn]["source_file"]}',
                           "issue": "duplicate road number across source files - first file kept"})
            continue

        acts = read_activities(wb[sheet_by_no[rn]], road["road_type"]) if rn in sheet_by_no else []
        acts.sort(key=order_key)
        present = {a["activity"] for a in acts if a["mapped"]}
        weights = config.road_weightage(present)
        for a in acts:
            a["weight"] = weights.get(a["activity"], 0)
            for fld, conv_fld in (("start_date", "start_date_converted"),
                                  ("end_date", "end_date_converted")):
                if a.get(conv_fld):
                    issues.append({"road_no": rn, "activity": a["raw_name"],
                                   "field": fld, "value": a.get(fld),
                                   "issue": "converted from dd/mm/yyyy text (confirmed by user)"})
                    continue
                prob = date_issue(a.get(fld))
                if prob and prob != "missing":
                    issues.append({"road_no": rn, "activity": a["raw_name"],
                                   "field": fld, "value": a.get(fld), "issue": prob})
            if a["start_date"] and a["end_date"] and ISO.match(a["start_date"]) \
                    and ISO.match(a["end_date"]) and a["end_date"] < a["start_date"]:
                issues.append({"road_no": rn, "activity": a["raw_name"],
                               "field": "range",
                               "value": f'{a["start_date"]}..{a["end_date"]}',
                               "issue": "end before start"})
            if not a["mapped"]:
                unmapped.setdefault(a["raw_name"], []).append(rn)
        road["activities"] = acts
        road["weightage"] = weights
        road["has_drain"] = road["drain_m"] > 0

        # Road-level start/end comes from Phase-Summary (confirmed: summary
        # is absolute truth). Flag, don't silently drop, when a sheet's own
        # activity rows span outside the summary's dates.
        sheet_starts = [a["start_date"] for a in acts
                        if a["start_date"] and ISO.match(a["start_date"])]
        sheet_ends = [a["end_date"] for a in acts
                      if a["end_date"] and ISO.match(a["end_date"])]
        if sheet_starts and min(sheet_starts) != road["start_date"]:
            issues.append({"road_no": rn, "activity": "(road sheet vs summary)",
                           "field": "start_date",
                           "value": f'sheet={min(sheet_starts)} summary={road["start_date"]}',
                           "issue": "sheet start differs from summary (summary used, per ruling)"})
        if sheet_ends and max(sheet_ends) != road["end_date"]:
            issues.append({"road_no": rn, "activity": "(road sheet vs summary)",
                           "field": "end_date",
                           "value": f'sheet={max(sheet_ends)} summary={road["end_date"]}',
                           "issue": "sheet end differs from summary (summary used, per ruling)"})
        roads_by_no[rn] = road


def build_dataset():
    """Parse every Source/*.xlsx and return the full dataset dict in memory.
    No disk writes -- the seed/verify scripts consume this return value."""
    files = source_files()
    roads_by_no, issues, unmapped = {}, [], {}
    if not files:
        issues.append({"road_no": "-", "activity": "(Source folder)", "field": "workbook",
                       "value": SRC_DIR, "issue": "no .xlsx file found in Source/"})
    for path in files:
        read_workbook(path, roads_by_no, issues, unmapped)

    roads = [roads_by_no[k] for k in sorted(roads_by_no)]
    total_len = sum(r["length_m"] for r in roads)
    for r in roads:
        r["length_weight_pct"] = round(r["length_m"] / total_len * 100, 2) if total_len else 0

    out = {
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "source_files": [os.path.basename(p) for p in files],
        "total_roads": len(roads),
        "total_length_m": total_len,
        "total_drain_m": sum(r["drain_m"] for r in roads),
        "total_cost_cr": round(sum((r["cost_cr"] or 0) for r in roads), 2),
        "activity_sequence": config.ACTIVITY_SEQUENCE,
        "major_weights": config.MAJOR_WEIGHTS,
        "roads": roads,
        "data_issues": issues,
        "unmapped_activities": [{"raw_name": k, "roads": v} for k, v in sorted(unmapped.items())],
    }
    return out


if __name__ == "__main__":
    d = build_dataset()
    print(f"files={len(d['source_files'])} roads={d['total_roads']} "
          f"total_len={d['total_length_m']} issues={len(d['data_issues'])} "
          f"unmapped_types={len(d['unmapped_activities'])}")
